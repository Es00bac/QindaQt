#!/usr/bin/env python3
"""Create a fictional, downloadable project for the QindaOffice showcase.

Requires python-docx and openpyxl. The same four documents are
opened by the installed QindaOffice apps for the site's captured walkthrough.
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets/sample"
INK = "172F3D"
JADE = "177B69"
ORANGE = "D27439"


def make_write() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.82)
    section.right_margin = Inches(.82)
    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)

    kicker = doc.add_paragraph("SIGNAL GARDEN  /  PROJECT BRIEF  /  01")
    kicker.runs[0].bold = True
    kicker.runs[0].font.size = Pt(8)
    kicker.runs[0].font.color.rgb = RGBColor.from_string(JADE)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Make the night feel alive.")
    run.bold = True
    run.font.size = Pt(23)
    run.font.color.rgb = RGBColor.from_string(INK)
    deck = doc.add_paragraph("A community light and sound evening · 14 November 2026")
    deck.runs[0].italic = True
    deck.runs[0].font.color.rgb = RGBColor.from_string(ORANGE)

    doc.add_heading("The idea", level=2)
    doc.add_paragraph(
        "Signal Garden turns an ordinary neighborhood courtyard into a one-night "
        "trail of light, sound, and small creative experiments. Guests follow four "
        "stations, meet local makers, and leave with a simple project to try at home."
    )
    doc.add_heading("What success looks like", level=2)
    for item in [
        "Welcome 320 guests across two timed entry windows.",
        "Keep the installation accessible, low-glare, and easy to navigate.",
        "Stay under the $8,600 budget ceiling and fund next year's workshop.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("The evening, at a glance", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for cell, label in zip(table.rows[0].cells, ["TIME", "MOMENT", "OWNER"]):
        cell.text = label
    for row in [
        ("17:30", "Doors and orientation", "Front of house"),
        ("18:00", "Light trail opens", "Installation crew"),
        ("19:15", "Maker talks and live demo", "Program team"),
        ("20:30", "Final walkthrough", "All teams"),
    ]:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = value
    doc.add_heading("Decision for Friday", level=2)
    doc.add_paragraph(
        "Confirm the accessible route and approve the weather cover before "
        "supplier deposits. The budget model and presentation use this same plan."
    )
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("SIGNAL GARDEN · WORKING DOCUMENT")
    doc.save(OUT / "signal-garden-brief.docx")


def make_calc() -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "SIGNAL GARDEN  /  BUDGET"
    sheet["A2"] = "One evening · 320 guests · numbers update from the assumptions below"
    headers = ["Cost line", "Qty", "Unit $", "Plan $", "Actual $", "Variance $"]
    for column, label in enumerate(headers, 1):
        sheet.cell(4, column, label)
    lines = [
        ("Lighting and power", 1, 2250, 2190),
        ("Artist fees", 4, 650, 2600),
        ("Access and wayfinding", 1, 850, 720),
        ("Site and weather cover", 1, 900, 980),
        ("Print and promotion", 1, 600, 470),
        ("Volunteer meals", 24, 18, 396),
    ]
    for index, (name, qty, price, actual) in enumerate(lines, 5):
        sheet.cell(index, 1, name)
        sheet.cell(index, 2, qty)
        sheet.cell(index, 3, price)
        sheet.cell(index, 4, f"=B{index}*C{index}")
        sheet.cell(index, 5, actual)
        sheet.cell(index, 6, f"=D{index}-E{index}")
    total = 11
    sheet.cell(total, 1, "TOTAL COST")
    for col in (4, 5, 6):
        letter = get_column_letter(col)
        sheet.cell(total, col, f"=SUM({letter}5:{letter}10)")
    sheet["A13"] = "REVENUE MODEL"
    sheet["A14"], sheet["B14"], sheet["C14"] = "Ticket", "Guests", "Revenue $"
    for row, label, count, value in [(15, "Standard", 240, 28), (16, "Pay-what-you-can", 80, 15)]:
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, count)
        sheet.cell(row, 3, f"=B{row}*{value}")
    sheet["A17"], sheet["C17"] = "Community sponsor", 2200
    sheet["A18"], sheet["C18"] = "TOTAL INCOME", "=SUM(C15:C17)"
    sheet["A19"], sheet["C19"] = "PROJECT BALANCE", "=C18-D11"
    sheet["A21"] = "Editable assumptions feed the formulas. Actuals are a rehearsal estimate."

    dark = PatternFill("solid", fgColor="CDECE3")
    mint = PatternFill("solid", fgColor="DCEFE6")
    orange = PatternFill("solid", fgColor="FBE8DA")
    for row in sheet.iter_rows(min_row=1, max_row=21, max_col=6):
        for cell in row:
            cell.font = Font(name="DejaVu Sans", size=10, color=INK)
            cell.alignment = Alignment(vertical="center")
            if cell.column > 1:
                cell.number_format = '#,##0;[Red](#,##0)'
    for row in (1, 4, 13, 14):
        for cell in sheet[row][:6 if row == 4 else 3]:
            cell.fill = dark
            cell.font = Font(name="DejaVu Sans", size=10, bold=True, color=INK)
    for row in (11, 18, 19):
        for cell in sheet[row][:6 if row == 11 else 3]:
            cell.fill = mint if row != 19 else orange
            cell.font = Font(name="DejaVu Sans", size=10, bold=True, color=INK)
    for row in range(5, 11):
        if row % 2 == 0:
            for cell in sheet[row]:
                cell.fill = PatternFill("solid", fgColor="F2F7F5")
    for col, width in {"A": 32, "B": 12, "C": 13, "D": 14, "E": 14, "F": 15}.items():
        sheet.column_dimensions[col].width = width
    sheet.row_dimensions[1].height = 30
    for row in (4, 13, 14, 18, 19):
        sheet.row_dimensions[row].height = 23
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    scenarios = book.create_sheet("Attendance")
    scenarios.sheet_view.showGridLines = False
    for row in [
        ["SIGNAL GARDEN  /  ATTENDANCE"],
        ["Case", "Guests", "Avg ticket $", "Ticket income $", "With sponsor $"],
        ["Quiet night", 220, 23, "=B3*C3", "=D3+Budget!C17"],
        ["Plan", 320, 25, "=B4*C4", "=D4+Budget!C17"],
        ["Full house", 400, 26, "=B5*C5", "=D5+Budget!C17"],
    ]:
        scenarios.append(row)
    for row in scenarios.iter_rows():
        for cell in row:
            cell.font = Font(name="DejaVu Sans", size=10, color=INK)
    for cell in scenarios[2]:
        cell.fill = dark
        cell.font = Font(name="DejaVu Sans", size=10, bold=True, color=INK)
    for col, width in {"A": 24, "B": 14, "C": 18, "D": 22, "E": 21}.items():
        scenarios.column_dimensions[col].width = width
    book.active = 0
    book.save(OUT / "signal-garden-budget.xlsx")
    chart = BarChart()
    chart.title = "Ticket income by attendance"
    chart.y_axis.title = "USD"
    chart.x_axis.title = "Scenario"
    chart.add_data(Reference(scenarios, min_col=4, min_row=2, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(scenarios, min_col=1, min_row=3, max_row=5))
    chart.width = 17
    chart.height = 8
    scenarios.add_chart(chart, "A8")
    book.save(OUT / "signal-garden-scenarios.xlsx")


def make_native_show() -> None:
    """A native FODP version preserves QindaShow's own slide styling."""
    ns = {
        "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
        "style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0",
        "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
        "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
        "svg": "urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0",
        "fo": "urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0",
        "presentation": "urn:oasis:names:tc:opendocument:xmlns:presentation:1.0",
        "qq": "https://qindaqt.example.org/qqo",
    }
    for prefix, uri in ns.items():
        ET.register_namespace(prefix, uri)

    def tag(prefix: str, name: str) -> str:
        return f"{{{ns[prefix]}}}{name}"

    root = ET.Element(tag("office", "document"), {tag("office", "mimetype"): "application/vnd.oasis.opendocument.presentation"})
    styles = ET.SubElement(root, tag("office", "automatic-styles"))
    layout = ET.SubElement(styles, tag("style", "page-layout"), {tag("style", "name"): "Default"})
    ET.SubElement(layout, tag("style", "page-layout-properties"), {
        tag("fo", "page-width"): "33.8667cm", tag("fo", "page-height"): "19.05cm",
    })
    masters = ET.SubElement(root, tag("office", "master-styles"))
    for id, bg, fg in [("night", "#172F3D", "#E8F4EC"), ("paper", "#E8F4EC", "#172F3D")]:
        ET.SubElement(masters, tag("style", "master-page"), {
            tag("style", "name"): id, tag("style", "page-layout-name"): "Default",
            tag("qq", "master-name"): id.title(), tag("qq", "background"): bg,
            tag("qq", "title-color"): fg, tag("qq", "body-color"): fg,
            tag("qq", "font"): "DejaVu Sans",
        })
    body = ET.SubElement(root, tag("office", "body"))
    presentation = ET.SubElement(body, tag("office", "presentation"))

    def add_style(id: str, fill: str | None = None) -> None:
        style = ET.SubElement(styles, tag("style", "style"), {
            tag("style", "name"): f"shape-{id}", tag("style", "family"): "graphic",
        })
        props = {tag("draw", "fill"): "solid" if fill else "none", tag("draw", "stroke"): "none"}
        if fill:
            props[tag("draw", "fill-color")] = fill
        ET.SubElement(style, tag("style", "graphic-properties"), props)

    def add_frame(page, id, x, y, w, h, value, size, color, bold=False, fill=None):
        add_style(id, fill)
        frame = ET.SubElement(page, tag("draw", "frame"), {
            tag("draw", "name"): id, tag("draw", "style-name"): f"shape-{id}",
            tag("svg", "x"): f"{x}pt", tag("svg", "y"): f"{y}pt",
            tag("svg", "width"): f"{w}pt", tag("svg", "height"): f"{h}pt",
        })
        box = ET.SubElement(frame, tag("draw", "text-box"))
        paragraph = ET.SubElement(box, tag("text", "p"))
        attrs = {tag("fo", "font-size"): f"{size}pt", tag("fo", "color"): color,
                 tag("style", "font-name"): "DejaVu Sans"}
        if bold:
            attrs[tag("fo", "font-weight")] = "bold"
        ET.SubElement(paragraph, tag("text", "span"), attrs).text = value

    slides = [
        ("The invitation", "Make the night", "feel alive.", "A community light and sound evening", "night", "Open with the courtyard as it is now, then describe the four stations that bring it to life."),
        ("The experience", "Four stops.", "One shared night.", "Light trail / Maker talks / Live demo / Take-home kit", "paper", "Walk the audience through timed entry, the accessible route, and where guests meet makers."),
        ("The plan", "320 guests.", "$8,600 ceiling.", "Timed entry keeps the route welcoming and the numbers clear.", "night", "The budget workbook currently projects 7,632 in costs and 10,120 in income."),
        ("The decision", "Build the route.", "Back the makers.", "Approve access and weather cover this Friday.", "paper", "Ask for approval of the step-free route and covered power before deposits are sent."),
    ]
    for n, (title, line1, line2, subtitle, master, speaker_notes) in enumerate(slides, 1):
        page = ET.SubElement(presentation, tag("draw", "page"), {
            tag("draw", "name"): f"slide{n}", tag("draw", "master-page-name"): master,
            tag("qq", "title"): title, tag("qq", "layout"): "title+content",
        })
        fg = "#E8F4EC" if master == "night" else "#172F3D"
        accent = "#91E0C7" if master == "night" else "#B9572C"
        add_frame(page, f"s{n}-rail", 46, 50, 10, 438, "", 1, fg, fill="#D27439")
        add_frame(page, f"s{n}-kicker", 85, 66, 740, 31, f"SIGNAL GARDEN  /  0{n}  /  {title.upper()}", 15, accent, True)
        add_frame(page, f"s{n}-title-a", 85, 164, 790, 66, line1, 45, fg, True)
        add_frame(page, f"s{n}-title-b", 85, 235, 790, 70, line2, 45, accent, True)
        add_frame(page, f"s{n}-detail", 87, 383, 760, 57, subtitle, 19, fg)
        add_frame(page, f"s{n}-folio", 835, 489, 100, 26, f"0{n} / 04", 12, accent, True)
        notes = ET.SubElement(page, tag("presentation", "notes"))
        notes_frame = ET.SubElement(notes, tag("draw", "frame"), {
            tag("presentation", "class"): "notes", tag("svg", "x"): "0pt",
            tag("svg", "y"): "0pt", tag("svg", "width"): "1pt",
            tag("svg", "height"): "1pt",
        })
        notes_box = ET.SubElement(notes_frame, tag("draw", "text-box"))
        ET.SubElement(notes_box, tag("text", "p")).text = speaker_notes
    (OUT / "signal-garden-pitch.fodp").write_bytes(ET.tostring(root, encoding="utf-8", xml_declaration=True))


def make_note() -> None:
    folder = OUT / "signal-garden-notes"
    page_dir = folder / "sections" / "fieldwork" / "pages"
    page_dir.mkdir(parents=True, exist_ok=True)
    (folder / "notebook.json").write_text(json.dumps({
        "name": "Signal Garden field notes",
        "sections": [{"id": "fieldwork", "name": "Fieldwork", "pages": [
            {"id": "walkthrough", "title": "Courtyard walkthrough"},
            {"id": "launch", "title": "Launch checklist"},
        ]}],
    }, indent=2) + "\n", encoding="utf-8")
    pages = {
        "walkthrough": [
            {"type": "text", "text": "COURTYARD WALKTHROUGH  /  18 SEPTEMBER\nThe entrance should feel like a small discovery, not a queue. Keep the first light station visible from the gate."},
            {"type": "check", "text": "Confirm step-free route and turning space", "checked": True},
            {"type": "check", "text": "Test low-glare markers after sunset", "checked": False},
            {"type": "check", "text": "Reserve covered power distribution", "checked": False},
            {"type": "text", "text": "Idea: invite each guest to leave one word on the closing wall. Those words become the opening of next year's brief."},
        ],
        "launch": [
            {"type": "text", "text": "LAUNCH CHECKLIST\nFriday is the go / no-go moment for supplier deposits."},
            {"type": "check", "text": "Sign off the accessible route", "checked": False},
            {"type": "check", "text": "Approve weather cover", "checked": False},
            {"type": "check", "text": "Send makers their run-of-show", "checked": False},
        ],
    }
    for id, blocks in pages.items():
        title = "Courtyard walkthrough" if id == "walkthrough" else "Launch checklist"
        (page_dir / f"{id}.json").write_text(json.dumps({
            "id": id, "title": title, "background": "dotted", "ink": [],
            "blocks": {"blocks": blocks},
        }, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    make_write()
    make_calc()
    make_native_show()
    make_note()
    included = [
        OUT / "signal-garden-brief.docx", OUT / "signal-garden-budget.xlsx",
        OUT / "signal-garden-scenarios.xlsx", OUT / "signal-garden-pitch.fodp",
        *sorted((OUT / "signal-garden-notes").rglob("*.json")),
    ]
    with zipfile.ZipFile(OUT / "signal-garden-project.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        for path in included:
            archive.write(path, path.relative_to(OUT))
    print(f"Wrote the working project and {OUT / 'signal-garden-project.zip'}")


if __name__ == "__main__":
    main()
